"""docxtpl engine: scan + deterministic render, formatting preserved."""

from io import BytesIO

from docx import Document

from neurolegal.templates.rendering.engine import default_fields, render_docx, scan_placeholders
from tests.unit.templates.helpers import build_docx, document_text

VALUES = {
    "landlord_fio": "Иванов Иван Иванович",
    "start_date": "01.09.2026",
    "rent_amount": "50000",
    "comment": "без животных",
}


def test_scan_finds_all_placeholders() -> None:
    names = scan_placeholders(build_docx())
    assert names == ["comment", "landlord_fio", "rent_amount", "start_date"]


def test_default_fields_are_text_required_with_name_as_label() -> None:
    fields = default_fields(["landlord_fio"])
    assert len(fields) == 1
    field = fields[0]
    assert (field.name, field.label, field.kind, field.required) == (
        "landlord_fio",
        "landlord_fio",
        "text",
        True,
    )
    assert field.hint is None


def test_render_substitutes_values_and_keeps_formatting() -> None:
    rendered = render_docx(build_docx(), VALUES)
    text = document_text(rendered)
    assert "Иванов Иван Иванович" in text
    assert "01.09.2026" in text
    assert "50000 руб./мес." in text
    assert "{{" not in text and "}}" not in text
    # Жирный run «Арендодатель: » пережил рендер
    doc = Document(BytesIO(rendered))
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert any("Арендодатель" in t for t in bold_runs)


def test_render_missing_optional_becomes_empty() -> None:
    values = {k: v for k, v in VALUES.items() if k != "comment"}
    text = document_text(render_docx(build_docx(), values))
    assert "Комментарий: \n" in text + "\n"
    assert "{{" not in text


def test_render_treats_user_values_as_data() -> None:
    values = VALUES | {"comment": "{{ hacked }} и <b>&"}
    text = document_text(render_docx(build_docx(), values))
    # Jinja-синтаксис и XML-спецсимволы в значении — литеральный текст
    assert "{{ hacked }} и <b>&" in text
