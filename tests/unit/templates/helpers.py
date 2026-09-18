"""Shared builder: a real .docx with {{ плейсхолдерами }} for engine/route tests."""

from io import BytesIO

from docx import Document


def build_docx() -> bytes:
    """Плейсхолдеры в разных местах: обычный абзац, жирный run рядом,
    ячейка таблицы — скан и рендер должны видеть все."""
    doc = Document()
    paragraph = doc.add_paragraph()
    label = paragraph.add_run("Арендодатель: ")
    label.bold = True
    paragraph.add_run("{{ landlord_fio }}")
    doc.add_paragraph("Дата начала: {{ start_date }}")
    doc.add_paragraph("Комментарий: {{ comment }}")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Плата: {{ rent_amount }} руб./мес."
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def document_text(data: bytes) -> str:
    """Весь текст документа, включая таблицы."""
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)
