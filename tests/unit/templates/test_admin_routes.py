"""Операторский CRUD /admin/templates* over SQLite + in-memory blob.

Гейт X-Internal-Token: no-op без настройки (локальная разработка), 401 при
неверном/отсутствующем токене, когда он задан.
"""

from collections.abc import AsyncIterator
from dataclasses import dataclass
from io import BytesIO

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from neurolegal.core.config import settings as core_settings
from neurolegal.templates.api.app import app
from neurolegal.templates.api.deps import get_blob_store, get_store
from neurolegal.templates.store.blob import InMemoryTemplatesBlobStore, template_key
from neurolegal.templates.store.models import Base, TplTemplate
from neurolegal.templates.store.template_store import TemplateStore
from tests.unit.templates.helpers import build_docx


def docx_with(*placeholders: str) -> bytes:
    doc = Document()
    for name in placeholders:
        doc.add_paragraph(f"Поле: {{{{ {name} }}}}")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


@dataclass
class Ctx:
    client: AsyncClient
    blob: InMemoryTemplatesBlobStore
    maker: async_sessionmaker[AsyncSession]


@pytest.fixture
async def ctx() -> AsyncIterator[Ctx]:
    engine = create_async_engine("sqlite+aiosqlite://")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, expire_on_commit=False)
    blob = InMemoryTemplatesBlobStore()

    async def _store() -> AsyncIterator[TemplateStore]:
        async with maker() as session:
            yield TemplateStore(session)

    app.dependency_overrides[get_store] = _store
    app.dependency_overrides[get_blob_store] = lambda: blob
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield Ctx(client=client, blob=blob, maker=maker)
    app.dependency_overrides.clear()
    await engine.dispose()


def _upload(name: str = "arenda.docx", data: bytes | None = None) -> dict[str, object]:
    return {"file": (name, data if data is not None else build_docx())}


CREATE_FORM = {
    "slug": "arenda-kvartiry",
    "title": "Аренда квартиры",
    "category": "Договоры",
    "description": "Договор аренды жилого помещения",
}


async def _create(ctx: Ctx) -> dict[str, object]:
    response = await ctx.client.post("/admin/templates", data=CREATE_FORM, files=_upload())
    assert response.status_code == 201, response.text
    body: dict[str, object] = response.json()
    return body


# --- гейт internal token ---


async def test_admin_routes_open_without_configured_token(ctx: Ctx) -> None:
    response = await ctx.client.get("/admin/templates")
    assert response.status_code == 200


async def test_admin_routes_reject_wrong_or_missing_token(
    ctx: Ctx, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(core_settings, "internal_token", "s3cret")
    for headers in ({}, {"X-Internal-Token": "wrong"}):
        response = await ctx.client.get("/admin/templates", headers=headers)
        assert response.status_code == 401, headers
    ok = await ctx.client.get("/admin/templates", headers={"X-Internal-Token": "s3cret"})
    assert ok.status_code == 200


async def test_public_routes_ignore_token_gate(ctx: Ctx, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(core_settings, "internal_token", "s3cret")
    response = await ctx.client.get("/templates")
    assert response.status_code == 200


# --- create ---


async def test_create_scans_fields_and_stores_docx(ctx: Ctx) -> None:
    body = await _create(ctx)
    assert body["slug"] == "arenda-kvartiry"
    assert body["status"] == "draft"
    fields = body["fields"]
    assert isinstance(fields, list)
    assert [f["name"] for f in fields] == ["comment", "landlord_fio", "rent_amount", "start_date"]
    assert all(f["kind"] == "text" and f["required"] for f in fields)
    assert template_key("arenda-kvartiry", "templates/") in ctx.blob.objects


async def test_create_duplicate_slug_is_409(ctx: Ctx) -> None:
    await _create(ctx)
    response = await ctx.client.post("/admin/templates", data=CREATE_FORM, files=_upload())
    assert response.status_code == 409


@pytest.mark.parametrize("slug", ["Аренда", "UPPER", "a b", "-lead", "a/../b"])
async def test_create_bad_slug_is_422(ctx: Ctx, slug: str) -> None:
    response = await ctx.client.post(
        "/admin/templates", data=CREATE_FORM | {"slug": slug}, files=_upload()
    )
    assert response.status_code == 422, slug


async def test_create_non_docx_and_broken_docx_are_422(ctx: Ctx) -> None:
    bad_name = await ctx.client.post(
        "/admin/templates", data=CREATE_FORM, files=_upload(name="t.pdf")
    )
    assert bad_name.status_code == 422
    broken = await ctx.client.post(
        "/admin/templates", data=CREATE_FORM, files=_upload(data=b"not a zip")
    )
    assert broken.status_code == 422


# --- list / draft visibility ---


async def test_admin_list_shows_draft_public_does_not(ctx: Ctx) -> None:
    await _create(ctx)
    admin_list = (await ctx.client.get("/admin/templates")).json()["templates"]
    assert [t["slug"] for t in admin_list] == ["arenda-kvartiry"]
    public_list = (await ctx.client.get("/templates")).json()["templates"]
    assert public_list == []


# --- patch ---


async def test_patch_publish_makes_template_public(ctx: Ctx) -> None:
    await _create(ctx)
    response = await ctx.client.patch(
        "/admin/templates/arenda-kvartiry", json={"status": "published"}
    )
    assert response.status_code == 200
    assert response.json()["status"] == "published"
    public_list = (await ctx.client.get("/templates")).json()["templates"]
    assert [t["slug"] for t in public_list] == ["arenda-kvartiry"]


async def test_patch_updates_metadata_and_fields(ctx: Ctx) -> None:
    created = await _create(ctx)
    fields = created["fields"]
    assert isinstance(fields, list)
    fields[0] = {**fields[0], "label": "Комментарий", "required": False}
    response = await ctx.client.patch(
        "/admin/templates/arenda-kvartiry",
        json={"title": "Аренда жилья", "fields": fields},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["title"] == "Аренда жилья"
    assert body["fields"][0]["label"] == "Комментарий"
    assert body["fields"][0]["required"] is False


async def test_patch_rejects_duplicate_field_names(ctx: Ctx) -> None:
    await _create(ctx)
    field = {"name": "x", "label": "X", "hint": None, "kind": "text", "required": True}
    response = await ctx.client.patch(
        "/admin/templates/arenda-kvartiry", json={"fields": [field, field]}
    )
    assert response.status_code == 422


async def test_patch_missing_template_is_404(ctx: Ctx) -> None:
    response = await ctx.client.patch("/admin/templates/missing", json={"title": "X"})
    assert response.status_code == 404


# --- replace file ---


async def test_replace_file_reports_added_and_orphaned_keeps_metadata(ctx: Ctx) -> None:
    await _create(ctx)
    relabel = await ctx.client.patch(
        "/admin/templates/arenda-kvartiry",
        json={
            "fields": [
                {"name": "landlord_fio", "label": "Арендодатель", "kind": "text", "required": True}
            ]
        },
    )
    assert relabel.status_code == 200
    new_docx = docx_with("landlord_fio", "tenant_fio")
    response = await ctx.client.put(
        "/admin/templates/arenda-kvartiry/file", files=_upload(data=new_docx)
    )
    assert response.status_code == 200
    body = response.json()
    assert body["added"] == ["tenant_fio"]
    assert body["orphaned"] == []
    by_name = {f["name"]: f for f in body["template"]["fields"]}
    # операторская подпись пережила замену файла
    assert by_name["landlord_fio"]["label"] == "Арендодатель"
    assert by_name["tenant_fio"]["label"] == "tenant_fio"
    # заменим файлом, где landlord_fio исчез — он станет осиротевшим, но останется
    only_tenant = docx_with("tenant_fio")
    second = await ctx.client.put(
        "/admin/templates/arenda-kvartiry/file", files=_upload(data=only_tenant)
    )
    assert second.status_code == 200
    assert second.json()["orphaned"] == ["landlord_fio"]
    assert {f["name"] for f in second.json()["template"]["fields"]} == {
        "landlord_fio",
        "tenant_fio",
    }


async def test_replace_file_missing_template_is_404(ctx: Ctx) -> None:
    response = await ctx.client.put("/admin/templates/missing/file", files=_upload())
    assert response.status_code == 404


# --- download / delete ---


async def test_download_returns_original_bytes(ctx: Ctx) -> None:
    await _create(ctx)
    response = await ctx.client.get("/admin/templates/arenda-kvartiry/file")
    assert response.status_code == 200
    assert response.content == ctx.blob.objects[template_key("arenda-kvartiry", "templates/")]
    assert 'filename="arenda-kvartiry.docx"' in response.headers["content-disposition"]


async def test_delete_removes_row_and_object(ctx: Ctx) -> None:
    await _create(ctx)
    response = await ctx.client.delete("/admin/templates/arenda-kvartiry")
    assert response.status_code == 204
    assert ctx.blob.objects == {}
    assert (await ctx.client.get("/admin/templates")).json()["templates"] == []


async def test_create_without_s3_is_503(ctx: Ctx) -> None:
    app.dependency_overrides[get_blob_store] = lambda: None
    response = await ctx.client.post("/admin/templates", data=CREATE_FORM, files=_upload())
    assert response.status_code == 503


async def test_delete_without_s3_still_removes_row(ctx: Ctx) -> None:
    await _create(ctx)
    app.dependency_overrides[get_blob_store] = lambda: None
    response = await ctx.client.delete("/admin/templates/arenda-kvartiry")
    assert response.status_code == 204
    assert (await ctx.client.get("/admin/templates")).json()["templates"] == []


async def test_delete_with_foreign_key_leaves_object_untouched(ctx: Ctx) -> None:
    # Строка, чей ключ мимо префикса: чужой объект не удаляется, строка — да
    async with ctx.maker() as session:
        session.add(
            TplTemplate(
                slug="evil",
                title="X",
                category="Y",
                description="",
                status="draft",
                s3_key="owner/u1/doc/original.docx",
                fields=[],
            )
        )
        await session.commit()
    await ctx.blob.put("owner/u1/doc/original.docx", b"user bytes")
    response = await ctx.client.delete("/admin/templates/evil")
    assert response.status_code == 204
    assert ctx.blob.objects == {"owner/u1/doc/original.docx": b"user bytes"}
