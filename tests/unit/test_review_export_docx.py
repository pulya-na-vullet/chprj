"""DOCX-экспорт отчёта (T-0049): читаем собранный файл обратно python-docx
и проверяем структуру — шапку, таблицу, порядок уровней, спецслучаи."""

from datetime import datetime
from io import BytesIO

from docx import Document

from neurolegal.agent.review.export import build_docx_report
from neurolegal.contracts import ReviewReportData
from neurolegal.contracts.chat import Citation
from neurolegal.contracts.review import ReviewCoverageItem, ReviewRisk

CREATED = datetime(2026, 7, 19, 14, 32)


def _risk(**over: object) -> ReviewRisk:
    base: dict[str, object] = {
        "rule_id": "r1",
        "title": "Неустойка",
        "level": "medium",
        "verdict": "confirmed",
        "section_number": "6.2",
        "contract_quote": "Пеня 0,5% в день",
        "explanation": "Не ограничена по сумме.",
        "recommendation": "Ограничить 10%.",
        "citations": [],
        "no_basis": False,
    }
    base.update(over)
    return ReviewRisk.model_validate(base)


def _report(**over: object) -> ReviewReportData:
    base: dict[str, object] = {
        "playbook_id": "supply_ru",
        "playbook_name": "Договор поставки",
        "document_id": "d1",
        "risks": [],
        "coverage": [],
        "disclaimer": "Черновая проверка: выводы требуют подтверждения юриста.",
        "role": "покупатель",
        "summary": "В целом рабочий договор.",
    }
    base.update(over)
    return ReviewReportData.model_validate(base)


def _texts(data: bytes) -> list[str]:
    doc = Document(BytesIO(data))
    return [p.text for p in doc.paragraphs if p.text]


def test_header_summary_counts_and_footer() -> None:
    report = _report(
        risks=[_risk(level="high")],
        coverage=[ReviewCoverageItem(rule_id="r1", title="Неустойка", status="risk")],
    )
    texts = _texts(build_docx_report(report, document_filename="dogovor.pdf", created_at=CREATED))

    assert "Проверка договора: Договор поставки" in texts
    assert "Документ: dogovor.pdf · вы — покупатель · 19.07.2026 14:32" in texts
    assert "Итог: В целом рабочий договор." in texts
    assert "Рисков: высокий 1 · средний 0 · низкий 0. Проверено 1 правил." in texts
    assert "Черновая проверка: выводы требуют подтверждения юриста." in texts
    assert "Сформировано в neurolegal · 19.07.2026 14:32" in texts


def test_header_omits_absent_parts() -> None:
    report = _report(role=None, summary=None)
    texts = _texts(build_docx_report(report, document_filename=None, created_at=CREATED))
    assert "Документ: 19.07.2026 14:32" in texts
    assert not any(t.startswith("Итог:") for t in texts)


def test_table_orders_by_level_and_renders_cells() -> None:
    citation = Citation(
        act_short_name="ГК РФ", kind="article", number="330", full_text="…", score=1.0
    )
    report = _report(
        risks=[
            _risk(rule_id="low1", level="low", title="Мелочь"),
            _risk(rule_id="high1", level="high", title="Крупный", citations=[citation]),
        ],
        coverage=[
            ReviewCoverageItem(rule_id="low1", title="Мелочь", status="risk"),
            ReviewCoverageItem(rule_id="high1", title="Крупный", status="risk"),
            ReviewCoverageItem(rule_id="ok1", title="Подсудность", status="ok"),
            ReviewCoverageItem(rule_id="na1", title="Экспорт", status="not_applicable"),
        ],
    )
    data = build_docx_report(report, document_filename=None, created_at=CREATED)
    doc = Document(BytesIO(data))
    table = doc.tables[0]

    headers = [c.text for c in table.rows[0].cells]
    assert headers == ["№ пункта", "Уровень", "Риск", "Рекомендация", "Основание"]
    # high прежде low, независимо от порядка в отчёте
    assert "Крупный" in table.rows[1].cells[2].text
    assert table.rows[1].cells[0].text == "п. 6.2"
    assert table.rows[1].cells[1].text == "Высокий"
    assert "Пеня 0,5% в день" in table.rows[1].cells[2].text
    assert table.rows[1].cells[3].text == "Ограничить 10%."
    assert table.rows[1].cells[4].text == "ст. 330 ГК РФ"
    assert "Мелочь" in table.rows[2].cells[2].text
    # хвост покрытия
    texts = _texts(data)
    assert "Остальные правила (1): без замечаний — Подсудность." in texts
    assert "Неприменимо: Экспорт." in texts


def test_missing_overstated_and_no_basis() -> None:
    report = _report(
        risks=[
            _risk(
                rule_id="m1", title="Приёмка", level="high", section_number=None, contract_quote=""
            ),
            _risk(rule_id="o1", title="Штраф", verdict="overstated", no_basis=True),
        ],
        coverage=[
            ReviewCoverageItem(rule_id="m1", title="Приёмка", status="missing"),
            ReviewCoverageItem(rule_id="o1", title="Штраф", status="risk"),
        ],
    )
    doc = Document(BytesIO(build_docx_report(report, document_filename=None, created_at=CREATED)))
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "—"
    assert "Приёмка — раздел в договоре не найден" in table.rows[1].cells[2].text
    assert "Штраф — возможно завышен" in table.rows[2].cells[2].text
    assert table.rows[2].cells[4].text == "норма в корпусе не найдена"


def test_no_risks_no_table() -> None:
    report = _report()
    doc = Document(BytesIO(build_docx_report(report, document_filename=None, created_at=CREATED)))
    assert doc.tables == []
