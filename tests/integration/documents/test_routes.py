import asyncio
import io

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from neurolegal.documents.api import deps, routes_documents
from neurolegal.documents.api.app import app
from neurolegal.documents.store.blob import InMemoryBlobStore
from neurolegal.documents.store.models import Base

pytestmark = pytest.mark.api_with_mocks


def _docx_bytes(paragraphs):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
async def client():
    engine = create_async_engine("sqlite+aiosqlite://")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    maker = async_sessionmaker(engine, expire_on_commit=False)
    blob = InMemoryBlobStore()

    async def _db_session():
        async with maker() as s:
            yield s

    app.dependency_overrides[deps.db_session] = _db_session
    app.dependency_overrides[deps.get_blob_store] = lambda: blob
    app.dependency_overrides[deps.get_session_factory] = lambda: maker
    transport = ASGITransport(app=app)
    # X-User-Id is a required header (T-0021) — send a fixed test user on
    # every request by default; tests that need a different owner override it.
    async with AsyncClient(
        transport=transport, base_url="http://test", headers={"X-User-Id": "u1"}
    ) as c:
        yield c, blob, maker
    app.dependency_overrides.clear()
    await engine.dispose()


async def test_upload_then_get_and_content(client):
    c, blob, _ = client
    files = {"file": ("c.docx", _docx_bytes(["1. Предмет", "текст"]), "application/octet-stream")}
    r = await c.post("/documents", files=files)
    assert r.status_code == 200, r.text
    doc_id = r.json()["id"]
    assert blob.objects  # original stored

    # The extraction runs fire-and-forget on the same loop. Polling GET alone
    # is a race against it on sqlite's single connection/worker-thread setup
    # (a tight sleep(0) poll loop can starve the background task of turns even
    # though it always completes) — await the tracked task directly for a
    # deterministic wait, then confirm the row reflects it.
    pending = list(routes_documents._background_tasks)
    if pending:
        await asyncio.gather(*pending)
    info = (await c.get(f"/documents/{doc_id}")).json()
    assert info["status"] == "ready"

    content = (await c.get(f"/documents/{doc_id}/content")).json()
    assert "Предмет" in content["full_text"]


async def test_upload_dedup_returns_same_id(client):
    c, *_ = client
    data = _docx_bytes(["1. Предмет"])
    files = {"file": ("c.docx", data, "application/octet-stream")}
    id1 = (await c.post("/documents", files=files)).json()["id"]
    files = {"file": ("c.docx", data, "application/octet-stream")}
    id2 = (await c.post("/documents", files=files)).json()["id"]
    assert id1 == id2


async def test_upload_dedup_does_not_cross_owners(client):
    """Одни и те же байты под другим владельцем — новая строка (T-0101, пункт 3).

    Дедуп ищет по (owner_id, content_hash). Если owner_id уйдёт из условия,
    второй пользователь получит в ответе чужой документ целиком: чужой id
    (и значит, доступ к /content и /download), чужое имя файла и чужую «Суть».
    """
    c, *_ = client
    data = _docx_bytes(["1. Предмет"])

    mine = (
        await c.post("/documents", files={"file": ("c.docx", data, "application/octet-stream")})
    ).json()
    assert mine["owner_id"] == "u1"

    theirs = (
        await c.post(
            "/documents",
            files={"file": ("c.docx", data, "application/octet-stream")},
            headers={"X-User-Id": "u2"},
        )
    ).json()
    assert theirs["owner_id"] == "u2"
    assert theirs["id"] != mine["id"]

    # Чужой документ и по прямой ссылке невидим...
    assert (await c.get(f"/documents/{theirs['id']}")).status_code == 404
    # ...и дедуп внутри своего владельца продолжает работать.
    again = (
        await c.post("/documents", files={"file": ("c.docx", data, "application/octet-stream")})
    ).json()
    assert again["id"] == mine["id"]


async def test_reject_unsupported_suffix(client):
    c, *_ = client
    files = {"file": ("x.txt", b"hello", "text/plain")}
    r = await c.post("/documents", files=files)
    assert r.status_code == 422


async def test_attach_and_list_for_conversation(client):
    c, *_ = client
    files = {"file": ("c.docx", _docx_bytes(["1. Предмет"]), "application/octet-stream")}
    doc_id = (await c.post("/documents", files=files)).json()["id"]
    r = await c.post(f"/documents/{doc_id}/attachments", json={"conversation_id": "conv-1"})
    assert r.status_code == 200
    listed = (await c.get("/conversations/conv-1/documents")).json()
    assert [d["id"] for d in listed["documents"]] == [doc_id]


async def test_download_redirects(client):
    c, *_ = client
    files = {"file": ("c.docx", _docx_bytes(["1. Предмет"]), "application/octet-stream")}
    doc_id = (await c.post("/documents", files=files)).json()["id"]
    r = await c.get(f"/documents/{doc_id}/download", follow_redirects=False)
    assert r.status_code == 307
    assert "memory://" in r.headers["location"]


async def test_upload_over_size_limit_rejected(client):
    c, *_ = client
    big = b"x" * (21 * 1024 * 1024)
    r = await c.post("/documents", files={"file": ("big.pdf", big, "application/pdf")})
    assert r.status_code == 422


async def test_download_gone_when_no_original(client):
    c, _blob, maker = client
    from neurolegal.documents.store.models import HubDocument

    async with maker() as s:
        row = HubDocument(
            owner_id="u1",
            filename="legacy.docx",
            content_type="application/octet-stream",
            size=0,
            content_hash="legacyhash",
            s3_key=None,
            status="ready",
            parser="docx",
            full_text="x",
            sections=[],
        )
        s.add(row)
        await s.commit()
        doc_id = row.id
    r = await c.get(f"/documents/{doc_id}/download", follow_redirects=False)
    assert r.status_code == 410


# --- C1: cross-tenant object authorization ---------------------------------
#
# The `client` fixture sends `X-User-Id: u1` by default. Every test below
# uploads a document as u1 (the owner) then reaches for it as u2 (the
# attacker) — none of these calls may succeed, and a wrong/missing owner must
# be indistinguishable from a document that doesn't exist at all (404, never
# a differentiated error).


async def _upload_as_u1(c) -> str:
    files = {"file": ("c.docx", _docx_bytes(["1. Предмет"]), "application/octet-stream")}
    r = await c.post("/documents", files=files)
    assert r.status_code == 200, r.text
    return str(r.json()["id"])


async def test_get_document_other_owner_is_404(client):
    c, *_ = client
    doc_id = await _upload_as_u1(c)
    r = await c.get(f"/documents/{doc_id}", headers={"X-User-Id": "u2"})
    assert r.status_code == 404
    # the owner can still fetch it
    r_owner = await c.get(f"/documents/{doc_id}")
    assert r_owner.status_code == 200


async def test_get_document_content_other_owner_is_404(client):
    c, *_ = client
    doc_id = await _upload_as_u1(c)
    r = await c.get(f"/documents/{doc_id}/content", headers={"X-User-Id": "u2"})
    assert r.status_code == 404


async def test_download_other_owner_is_404(client):
    c, *_ = client
    doc_id = await _upload_as_u1(c)
    r = await c.get(
        f"/documents/{doc_id}/download", headers={"X-User-Id": "u2"}, follow_redirects=False
    )
    assert r.status_code == 404
    r_owner = await c.get(f"/documents/{doc_id}/download", follow_redirects=False)
    assert r_owner.status_code == 307


async def test_delete_other_owner_is_404_and_does_not_delete(client):
    c, blob, _ = client
    doc_id = await _upload_as_u1(c)
    r = await c.delete(f"/documents/{doc_id}", headers={"X-User-Id": "u2"})
    assert r.status_code == 404
    assert blob.objects  # the original is still there — nothing was deleted
    r_owner = await c.get(f"/documents/{doc_id}")
    assert r_owner.status_code == 200


async def test_attach_other_owner_is_404(client):
    c, *_ = client
    doc_id = await _upload_as_u1(c)
    r = await c.post(
        f"/documents/{doc_id}/attachments",
        json={"conversation_id": "conv-x"},
        headers={"X-User-Id": "u2"},
    )
    assert r.status_code == 404
    listed = (await c.get("/conversations/conv-x/documents")).json()
    assert listed["documents"] == []


async def test_detach_other_owner_is_404(client):
    c, *_ = client
    doc_id = await _upload_as_u1(c)
    await c.post(f"/documents/{doc_id}/attachments", json={"conversation_id": "conv-y"})
    r = await c.delete(f"/documents/{doc_id}/attachments/conv-y", headers={"X-User-Id": "u2"})
    assert r.status_code == 404
    # attachment survives the attacker's detach attempt
    listed = (await c.get("/conversations/conv-y/documents")).json()
    assert [d["id"] for d in listed["documents"]] == [doc_id]


async def test_conversation_documents_excludes_other_owner(client):
    c, *_ = client
    doc_id = await _upload_as_u1(c)
    await c.post(f"/documents/{doc_id}/attachments", json={"conversation_id": "conv-z"})
    listed = (await c.get("/conversations/conv-z/documents", headers={"X-User-Id": "u2"})).json()
    assert listed["documents"] == []
    listed_owner = (await c.get("/conversations/conv-z/documents")).json()
    assert [d["id"] for d in listed_owner["documents"]] == [doc_id]


async def test_missing_x_user_id_is_422_on_per_document_routes(client):
    c, *_ = client
    doc_id = await _upload_as_u1(c)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as bare:
        assert (await bare.get(f"/documents/{doc_id}")).status_code == 422
        assert (await bare.get(f"/documents/{doc_id}/content")).status_code == 422
        assert (
            await bare.get(f"/documents/{doc_id}/download", follow_redirects=False)
        ).status_code == 422
        assert (await bare.delete(f"/documents/{doc_id}")).status_code == 422
        assert (
            await bare.post(f"/documents/{doc_id}/attachments", json={"conversation_id": "conv-w"})
        ).status_code == 422
        assert (await bare.delete(f"/documents/{doc_id}/attachments/conv-w")).status_code == 422
        assert (await bare.get("/conversations/conv-w/documents")).status_code == 422
