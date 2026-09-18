"""DOCX-экспорт отчёта по рискам (T-0049).

Повторяет утверждённый текст-формат §6 спеки E07, но риски идут таблицей
«№ пункта | Уровень | Риск | Рекомендация | Основание» (спека §3). Сборка
детерминированная, без LLM; зеркалит правила frontend/src/util/reviewExport.ts
(missing-правило, «возможно завышен», «норма в корпусе не найдена»)."""

from datetime import datetime
from io import BytesIO

from docx import Document

from neurolegal.contracts import ReviewReportData
from neurolegal.contracts.review import ReviewRisk

_LEVEL_ORDER: tuple[str, ...] = ("high", "medium", "low")
_LEVEL_LABELS = {"high": "Высокий", "medium": "Средний", "low": "Низкий"}
_TABLE_HEADERS = ("№ пункта", "Уровень", "Риск", "Рекомендация", "Основание")


def _basis_text(risk: ReviewRisk) -> str:
    if risk.no_basis or not risk.citations:
        return "норма в корпусе не найдена"
    return ", ".join(f"ст. {c.number} {c.act_short_name}" for c in risk.citations)


def build_docx_report(
    report: ReviewReportData,
    *,
    document_filename: str | None,
    created_at: datetime,
) -> bytes:
    doc = Document()
    doc.add_heading(f"Проверка договора: {report.playbook_name}", level=1)

    when = created_at.strftime("%d.%m.%Y %H:%M")
    role = f"вы — {report.role}" if report.role else None
    meta = [part for part in (document_filename, role, when) if part]
    doc.add_paragraph("Документ: " + " · ".join(meta))

    if report.summary:
        doc.add_paragraph(f"Итог: {report.summary}")

    counts = {level: sum(1 for r in report.risks if r.level == level) for level in _LEVEL_ORDER}
    doc.add_paragraph(
        f"Рисков: высокий {counts['high']} · средний {counts['medium']} · "
        f"низкий {counts['low']}. Проверено {len(report.coverage)} правил."
    )

    missing_rule_ids = {c.rule_id for c in report.coverage if c.status == "missing"}
    ordered = [r for level in _LEVEL_ORDER for r in report.risks if r.level == level]
    if ordered:
        table = doc.add_table(rows=1, cols=len(_TABLE_HEADERS))
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, _TABLE_HEADERS, strict=True):
            cell.paragraphs[0].add_run(header).bold = True
        for risk in ordered:
            cells = table.add_row().cells
            is_missing = risk.rule_id in missing_rule_ids
            has_section = not is_missing and risk.section_number
            cells[0].text = f"п. {risk.section_number}" if has_section else "—"
            cells[1].text = _LEVEL_LABELS[risk.level]
            title = risk.title
            if is_missing:
                title += " — раздел в договоре не найден"
            elif risk.verdict == "overstated":
                title += " — возможно завышен"
            cells[2].paragraphs[0].add_run(title).bold = True
            if not is_missing and risk.contract_quote:
                quote = cells[2].add_paragraph()
                quote.add_run(f"«{risk.contract_quote}»").italic = True
            if risk.explanation:
                cells[2].add_paragraph(risk.explanation)
            cells[3].text = risk.recommendation
            cells[4].text = _basis_text(risk)

    ok = [c.title for c in report.coverage if c.status == "ok"]
    not_applicable = [c.title for c in report.coverage if c.status == "not_applicable"]
    if ok:
        doc.add_paragraph(f"Остальные правила ({len(ok)}): без замечаний — {', '.join(ok)}.")
    if not_applicable:
        doc.add_paragraph(f"Неприменимо: {', '.join(not_applicable)}.")

    doc.add_paragraph(report.disclaimer)
    doc.add_paragraph(f"Сформировано в neurolegal · {when}")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
